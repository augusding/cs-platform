"""
文档解析模块。
支持 PDF / Excel / Word，返回纯文本列表（按页/按sheet）。
"""
import logging
from pathlib import Path

logger = logging.getLogger(__name__)


def parse_pdf(file_path: str) -> list[str]:
    """PDF → 每页文本列表"""
    import pdfplumber
    pages: list[str] = []
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            if text.strip():
                pages.append(text.strip())
    return pages


def parse_excel(file_path: str) -> list[str]:
    """Excel → 每个 sheet 转为文本（表头+行）"""
    import openpyxl
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    sheets: list[str] = []
    for sheet in wb.worksheets:
        rows: list[str] = []
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            if any(c.strip() for c in cells):
                rows.append(" | ".join(cells))
        if rows:
            sheets.append(f"[{sheet.title}]\n" + "\n".join(rows))
    wb.close()
    return sheets


def parse_word(file_path: str) -> list[str]:
    """Word → 段落列表（过滤空段落）"""
    from docx import Document
    doc = Document(file_path)
    paragraphs: list[str] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            paragraphs.append(text)
    return paragraphs


def parse_text(file_path: str) -> list[str]:
    """纯文本文件 → 全文作为单页。尝试 utf-8，失败回退 gbk。"""
    for encoding in ("utf-8", "gbk", "latin-1"):
        try:
            with open(file_path, "r", encoding=encoding) as f:
                content = f.read().strip()
            if content:
                return [content]
            return []
        except UnicodeDecodeError:
            continue
    return []


def parse_csv(file_path: str) -> list[str]:
    """CSV → 每行转为 'col1 | col2 | col3' 文本"""
    import csv
    rows: list[str] = []
    for encoding in ("utf-8", "gbk", "latin-1"):
        try:
            with open(file_path, "r", encoding=encoding, newline="") as f:
                reader = csv.reader(f)
                for row in reader:
                    cells = [str(c).strip() for c in row]
                    if any(cells):
                        rows.append(" | ".join(cells))
            break
        except UnicodeDecodeError:
            rows = []
            continue
    return ["\n".join(rows)] if rows else []


def parse_file(file_path: str) -> list[str]:
    """根据扩展名自动选择解析器"""
    suffix = Path(file_path).suffix.lower()
    if suffix == ".pdf":
        return parse_pdf(file_path)
    if suffix in (".xlsx", ".xls"):
        return parse_excel(file_path)
    if suffix in (".docx", ".doc"):
        return parse_word(file_path)
    if suffix == ".txt":
        return parse_text(file_path)
    if suffix == ".csv":
        return parse_csv(file_path)
    raise ValueError(f"Unsupported file type: {suffix}")
